#!/usr/bin/env python3
"""The lead tracker (D1 default: a Google Sheet "Boost2Business · לידים").

Builds an EMPTY .xlsx with the structure from docs/TASK.md §1.1. Upload it to
Drive with conversion and it becomes a Google Sheet; dropdowns, formulas and
conditional formatting survive the conversion.

Tabs:
  לידים   one row per lead; columns = data/leads.example.csv (Hebrew headers,
          the English key in each header's note, for n8n later); dropdowns;
          overdue next_action_date in red, due today in amber.
  משפך    funnel by formulas: last 12 weeks (by lead date) and by source.
  הסבר    how to use it, 5 lines.
  רשימות  the dropdown values. Edit here, and the dropdowns follow.

No lead rows are written: real lead data never goes to git (CLAUDE.md).

usage:  python3 tools/make_tracker.py [out.xlsx]     (default: ops/lead-tracker.xlsx)
"""
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROWS = 500  # data rows 2..ROWS+1
LAST = ROWS + 1

# key -> (Hebrew header, width)
HEAD = {
    "lead_id": ("מס׳ ליד", 9), "created_at": ("תאריך פנייה", 12), "name": ("שם", 16), "phone": ("טלפון", 13),
    "source": ("מקור", 15), "track": ("מסלול", 10), "first_message": ("ההודעה הראשונה (בקצרה)", 30),
    "stage": ("שלב", 16), "next_action": ("הצעד הבא", 22), "next_action_date": ("עד מתי", 11),
    "call_date": ("תאריך שיחה", 11), "call_summary_link": ("קישור לסיכום שיחה", 16), "quote_id": ("מס׳ הצעה", 11),
    "quote_amount": ("סכום הצעה ₪", 11), "quote_status": ("סטטוס הצעה", 12), "paid_status": ("תשלום", 11),
    "project_status": ("סטטוס פרויקט", 12), "hours_spent": ("השעות שלי", 9), "lost_reason": ("למה לא המשיך", 16),
    "consent_to_contact": ("הסכים שאחזור אליו", 10), "notes": ("הערות", 28),
}

# the stage values from docs/PIPELINE.md, in pipeline order (the order is the rank used by the funnel)
LISTS = {
    "stage": ["0 ליד חדש", "1 בסינון", "2 שיחה נקבעה", "2b שיחה התקיימה", "3 הצעה נשלחה", "4 ממתין לתשלום",
              "5 באפיון מלא", "6 ביישום", "8 באוויר", "9 בתמיכה", "X לא רלוונטי", "Z בהמתנה"],
    "source": ["אתר/דף הבית", "אתר/פתרונות", "אתר/למידה", "אתר/עמוד סוכן", "לינקדאין", "פייסבוק",
               "וואטסאפ ישיר", "המלצה מפה לאוזן", "אחר", "לא ידוע"],
    "track": ["פתרון", "סוכן", "למידה", "עוד לא ברור"],
    "lost_reason": ["לא מתאים לשירות", "תקציב", "לא עכשיו / תזמון", "לא ענה", "בחר ספק אחר", "עושה לבד", "אחר"],
    "quote_status": ["טיוטה", "נשלחה", "אושרה", "נדחתה", "פג תוקף"],
    "paid_status": ["עוד לא", "ממתין", "שולם חלקית", "שולם"],
    "project_status": ["לא התחיל", "באפיון", "ביישום", "באוויר", "בתמיכה", "הסתיים"],
    "consent_to_contact": ["כן", "לא"],
}

NIGHT = "080A24"
INDIGO = "4F46E5"
HEAD_FILL = PatternFill("solid", fgColor="1E1B4B")
HEAD_FONT = Font(name="Rubik", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Rubik", size=11)
THIN = Side(style="thin", color="D9DAF2")
GRID = Border(bottom=THIN, left=THIN, right=THIN, top=THIN)


def keys():
    with open(os.path.join(ROOT, "data", "leads.example.csv"), encoding="utf-8") as f:
        return next(csv.reader(f))


def build(out):
    cols = keys()
    missing = [k for k in cols if k not in HEAD]
    if missing:
        raise SystemExit(f"no Hebrew header for: {missing} (update HEAD in this file)")
    col = {k: get_column_letter(i) for i, k in enumerate(cols, 1)}
    rank_col = get_column_letter(len(cols) + 1)  # hidden helper: stage rank

    wb = Workbook()

    # ---------- רשימות ----------
    ls = wb.active
    ls.title = "רשימות"
    ls.sheet_view.rightToLeft = True
    list_ref = {}
    for i, (k, vals) in enumerate(LISTS.items(), 1):
        c = get_column_letter(i)
        ls[f"{c}1"] = HEAD[k][0]
        ls[f"{c}1"].font = HEAD_FONT
        ls[f"{c}1"].fill = HEAD_FILL
        for r, v in enumerate(vals, 2):
            ls[f"{c}{r}"] = v
            ls[f"{c}{r}"].font = BODY_FONT
        # room for Karin to add values: the dropdown range runs to row 30
        list_ref[k] = f"'רשימות'!${c}$2:${c}$30"
        ls.column_dimensions[c].width = 20
    ls["A32"] = "אפשר להוסיף ערכים בכל עמודה (עד שורה 30). את סדר השלבים בעמודה הראשונה לא משנים: המשפך נשען עליו."
    ls["A32"].font = Font(name="Rubik", italic=True, color="6B6F9E")

    # ---------- לידים ----------
    ws = wb.create_sheet("לידים", 0)
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "D2"
    for k in cols:
        c = col[k]
        cell = ws[f"{c}1"]
        cell.value = HEAD[k][0]
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.comment = Comment(f"שם השדה: {k}", "Boost2Business")
        ws.column_dimensions[c].width = HEAD[k][1]
    ws.row_dimensions[1].height = 34
    ws[f"{rank_col}1"] = "stage_rank"
    ws.column_dimensions[rank_col].hidden = True
    ws.auto_filter.ref = f"A1:{col[cols[-1]]}{LAST}"

    stage_rng = f"'רשימות'!$A$2:$A$13"
    for r in range(2, LAST + 1):
        ws[f"{col['lead_id']}{r}"] = f'=IF({col["name"]}{r}="","","L-"&TEXT(ROW()-1,"0000"))'
        # 1..12 in pipeline order, 0 for an empty row (never "", so numeric comparisons stay safe)
        ws[f"{rank_col}{r}"] = f'=IF({col["stage"]}{r}="",0,IFERROR(MATCH({col["stage"]}{r},{stage_rng},0),0))'
        for k in ("created_at", "next_action_date", "call_date"):
            ws[f"{col[k]}{r}"].number_format = "dd/mm/yyyy"
        ws[f"{col['quote_amount']}{r}"].number_format = "#,##0"
        ws[f"{col['phone']}{r}"].number_format = "@"

    for k, rng in list_ref.items():
        if k == "stage":
            rng = stage_rng
        dv = DataValidation(type="list", formula1=f"={rng}", allow_blank=True, showErrorMessage=k == "stage",
                            errorTitle="שלב לא מוכר", error="בחרי שלב מהרשימה")
        dv.add(f"{col[k]}2:{col[k]}{LAST}")
        ws.add_data_validation(dv)
    dvd = DataValidation(type="date", operator="greaterThan", formula1="DATE(2025,1,1)", allow_blank=True)
    for k in ("created_at", "next_action_date", "call_date"):
        dvd.add(f"{col[k]}2:{col[k]}{LAST}")
    ws.add_data_validation(dvd)

    # overdue / due today, only for leads that are still open (not X, not Z, not live or in support)
    nad, st = col["next_action_date"], col["stage"]
    open_ = f'AND(${st}2<>"",LEFT(${st}2,1)<>"X",LEFT(${st}2,1)<>"Z",LEFT(${st}2,1)<>"8",LEFT(${st}2,1)<>"9")'
    full = f"A2:{col[cols[-1]]}{LAST}"
    ws.conditional_formatting.add(f"{nad}2:{nad}{LAST}", FormulaRule(
        formula=[f'AND(${nad}2<>"",${nad}2<TODAY(),{open_})'],
        fill=PatternFill("solid", fgColor="EF4444"), font=Font(color="FFFFFF", bold=True), stopIfTrue=True))
    ws.conditional_formatting.add(full, FormulaRule(
        formula=[f'AND(${nad}2<>"",${nad}2<TODAY(),{open_})'], fill=PatternFill("solid", fgColor="FEE2E2")))
    ws.conditional_formatting.add(f"{nad}2:{nad}{LAST}", FormulaRule(
        formula=[f'AND(${nad}2<>"",${nad}2=TODAY(),{open_})'],
        fill=PatternFill("solid", fgColor="FBBF24"), font=Font(bold=True)))
    ws.conditional_formatting.add(full, FormulaRule(
        formula=[f'LEFT(${st}2,1)="X"'], font=Font(color="9CA3AF")))

    # ---------- משפך ----------
    fs = wb.create_sheet("משפך", 1)
    fs.sheet_view.rightToLeft = True
    L = lambda k: f"'לידים'!${col[k]}$2:${col[k]}${LAST}"  # noqa: E731
    R = f"'לידים'!${rank_col}$2:${rank_col}${LAST}"
    has = f'({L("name")}<>"")'
    # "reached" conditions. The tracker keeps only the current stage, so a stage counts when the lead is
    # at it or past it (ranks 1..10; X=11 and Z=12 are not "past"), or when the column that stage fills is set.
    reached = {
        "לידים": has,
        "שיחות נקבעו": f'{has}*((({L("call_date")}<>"")+({R}>=3)*({R}<=10))>0)',
        "שיחות התקיימו": f'{has}*((({R}>=4)*({R}<=10)+({L("quote_id")}<>""))>0)',
        "הצעות נשלחו": f'{has}*((({R}>=5)*({R}<=10)+({L("quote_id")}<>""))>0)',
        "נסגרו (אושרה)": f'{has}*((({R}>=6)*({R}<=10)+({L("quote_status")}="אושרה"))>0)',
        "שולמו": f'{has}*((({R}>=7)*({R}<=10)+({L("paid_status")}="שולם"))>0)',
        "באוויר": f'{has}*({R}>=9)*({R}<=10)',
    }
    names = list(reached)
    title_font = Font(name="Secular One", size=16, color=INDIGO)
    fs["A1"] = "משפך הלידים · מחושב לבד מהטבלה"
    fs["A1"].font = title_font
    fs["A2"] = "כל ליד נספר בשלב שבו הוא נמצא ובכל השלבים שלפניו. השבוע נקבע לפי תאריך הפנייה."
    fs["A2"].font = Font(name="Rubik", color="6B6F9E")

    def header(row, first):
        for i, h in enumerate([first] + names, 1):
            c = fs.cell(row=row, column=i, value=h)
            c.font, c.fill = HEAD_FONT, HEAD_FILL
            c.alignment = Alignment(horizontal="center", wrap_text=True)

    # by week (Sunday start), last 12 weeks, newest on top
    fs["A4"] = "לפי שבוע"
    fs["A4"].font = Font(name="Rubik", bold=True, size=13)
    header(5, "שבוע שמתחיל ב-")
    D = L("created_at")
    for n in range(12):
        r = 6 + n
        fs.cell(row=r, column=1, value=f"=TODAY()-WEEKDAY(TODAY())+1-7*{n}").number_format = "dd/mm/yyyy"
        for i, h in enumerate(names, 2):
            fs.cell(row=r, column=i, value=f"=SUMPRODUCT({reached[h]}*({D}>=$A{r})*({D}<$A{r}+7))")
    tot = 18
    fs.cell(row=tot, column=1, value="סה״כ מאז ההתחלה").font = Font(bold=True)
    for i, h in enumerate(names, 2):
        fs.cell(row=tot, column=i, value=f"=SUMPRODUCT({reached[h]})").font = Font(bold=True)
    fs.cell(row=tot + 1, column=1, value="אחוז מהלידים").font = Font(color="6B6F9E")
    for i in range(2, len(names) + 2):
        c = get_column_letter(i)
        fs.cell(row=tot + 1, column=i, value=f"=IF($B${tot}=0,\"\",{c}{tot}/$B${tot})").number_format = "0%"

    # by source
    s0 = 22
    fs.cell(row=s0, column=1, value="לפי מקור (מאז ההתחלה)").font = Font(name="Rubik", bold=True, size=13)
    header(s0 + 1, "מקור")
    for j, src in enumerate(LISTS["source"]):
        r = s0 + 2 + j
        fs.cell(row=r, column=1, value=f"='רשימות'!$B${2 + j}")
        for i, h in enumerate(names, 2):
            fs.cell(row=r, column=i, value=f"=SUMPRODUCT({reached[h]}*({L('source')}=$A{r}))")
    e = s0 + 2 + len(LISTS["source"]) + 1
    fs.cell(row=e, column=1, value="הערה: כל עוד הקישור לוואטסאפ באתר בלי טקסט מוכן (D2), רוב הלידים יהיו \"לא ידוע\".").font = Font(color="6B6F9E", italic=True)

    k0 = e + 2
    fs.cell(row=k0, column=1, value="השעות שלי").font = Font(name="Rubik", bold=True, size=13)
    fs.cell(row=k0 + 1, column=1, value="סה״כ שעות על לידים ולקוחות")
    fs.cell(row=k0 + 1, column=2, value=f"=SUM({L('hours_spent')})").number_format = "0.0"
    fs.cell(row=k0 + 2, column=1, value="ממוצע שעות ללקוח שסגר")
    fs.cell(row=k0 + 2, column=2, value=f'=IFERROR(SUMPRODUCT(({R}>=6)*({R}<=10),{L("hours_spent")})/SUMPRODUCT(({R}>=6)*({R}<=10)),"")').number_format = "0.0"
    fs.cell(row=k0 + 3, column=1, value="ממוצע שעות לליד שלא נסגר")
    fs.cell(row=k0 + 3, column=2, value=f'=IFERROR(SUMPRODUCT(({R}=11),{L("hours_spent")})/SUMPRODUCT(({R}=11)*1),"")').number_format = "0.0"
    fs.cell(row=k0 + 4, column=1, value="ימים ממוצעים בכל שלב: יגיע בשלב 3 (אוטומציה ב-n8n שרושמת תאריך לכל מעבר שלב).").font = Font(color="6B6F9E", italic=True)

    fs.column_dimensions["A"].width = 26
    for i in range(2, len(names) + 2):
        fs.column_dimensions[get_column_letter(i)].width = 13
    for row in fs.iter_rows(min_row=5, max_row=e, max_col=len(names) + 1):
        for c in row:
            if c.value is not None:
                c.border = GRID

    # ---------- הסבר ----------
    hs = wb.create_sheet("הסבר", 2)
    hs.sheet_view.rightToLeft = True
    hs["A1"] = "איך עובדים עם הטבלה · 5 שורות"
    hs["A1"].font = title_font
    lines = [
        "1. כל פנייה בוואטסאפ = שורה חדשה בלשונית \"לידים\": תאריך, שם, טלפון, מקור, ההודעה בקצרה. מס׳ הליד נוצר לבד.",
        "2. בכל פעם שמשהו זז, מעדכנים 3 תאים: שלב, הצעד הבא, ועד מתי. אדום = עבר התאריך, כתום = היום.",
        "3. אחרי שיחת אפיון: תאריך שיחה, קישור לסיכום, ושלב \"2b\". אחרי שליחת הצעה: מס׳ הצעה וסטטוס.",
        "4. ליד שלא ממשיך: שלב \"X לא רלוונטי\" וסיבה. משאירים את מס׳ ההצעה אם הייתה, כדי שהמשפך יספור נכון.",
        "5. לשונית \"משפך\" מתעדכנת לבד. בלשונית \"רשימות\" אפשר להוסיף מקורות וסיבות. רק מה שצריך: שם, טלפון והצורך.",
    ]
    for i, t in enumerate(lines, 3):
        hs[f"A{i}"] = t
        hs[f"A{i}"].font = Font(name="Rubik", size=12)
        hs[f"A{i}"].alignment = Alignment(wrap_text=True, vertical="top")
        hs.row_dimensions[i].height = 36
    hs.column_dimensions["A"].width = 110
    hs["A10"] = "פרטיות: זה מידע אישי של לקוחות. לא משתפים את הטבלה, לא מעתיקים אותה ל-GitHub, ושומרים רק מה שהשלב צריך."
    hs["A10"].font = Font(name="Rubik", size=11, color="B91C1C", bold=True)
    hs.sheet_properties.tabColor = INDIGO
    ws.sheet_properties.tabColor = "67E8F9"
    fs.sheet_properties.tabColor = "8B5CF6"
    ls.sheet_properties.tabColor = "A9AEE0"

    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    wb.save(out)
    print(f"wrote {out}: {len(cols)} columns, {ROWS} empty rows, tabs {wb.sheetnames}")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "ops", "lead-tracker.xlsx"))
